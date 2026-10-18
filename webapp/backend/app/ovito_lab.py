"""OVITO page API: runs the official OVITO Python module (pip package `ovito`) in one long-lived worker process.

All analysis and rendering on this page is done by OVITO itself; this module only starts the worker, passes it the
page settings, and serves the images and exports it writes into a temporary folder (never into the project or the
data folders). Long operations run as background jobs.
"""
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import uuid
from pathlib import Path

import numpy as np
from fastapi import APIRouter, Body, HTTPException, Request
from fastapi.responses import FileResponse, Response

from . import blobstore, jobs, ovito_script, remote, studio
from .settings import HOSTED

router = APIRouter(prefix="/api/ovito", tags=["ovito"])
WORKER = Path(__file__).resolve().parents[1] / "ovito_worker.py"
OUT_DIR = Path(tempfile.gettempdir()) / "nanocut_ovito"
OUT_DIR.mkdir(exist_ok=True)
UPLOAD_DIR = OUT_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)
UPLOAD_KEEP_S = 24 * 3600
KEEP_FILES = 60
TIMEOUT_S = float(os.environ.get("NANOCUT_OVITO_TIMEOUT", "1800"))
EXPORT_FORMATS = {"lammps/dump": ".dump", "xyz": ".xyz", "lammps/data": ".data", "vasp": ".poscar"}


class _Worker:
    """One OVITO process; requests are serialised (OVITO keeps the loaded frame cached between them)."""

    def __init__(self):
        self.proc, self.lock, self.info = None, threading.Lock(), None

    def _start(self):
        port, token = remote.bridge()                            # HPC frames reach the worker through this, in memory
        self.proc = subprocess.Popen([sys.executable, str(WORKER)], stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                     stderr=subprocess.DEVNULL, text=True, encoding="utf-8",
                                     env={**os.environ, "NANOCUT_BRIDGE": f"{port}:{token}"},
                                     creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        self.info = self._send({"op": "hello"}, 120)

    def _send(self, req, timeout):
        self.proc.stdin.write(json.dumps(req) + "\n")
        self.proc.stdin.flush()
        box = {}
        t = threading.Thread(target=lambda: box.setdefault("line", self.proc.stdout.readline()), daemon=True)
        t.start()
        t.join(timeout)
        if "line" not in box or not box["line"]:
            self.stop()
            raise RuntimeError("the OVITO worker did not answer (timed out or crashed); it will be restarted on the next request")
        res = json.loads(box["line"])
        if not res.get("ok"):
            raise RuntimeError(res.get("error", "OVITO error"))
        return res["result"]

    def ask(self, req, timeout=TIMEOUT_S):
        with self.lock:
            if self.proc is None or self.proc.poll() is not None:
                self._start()
            return self._send(req, timeout)

    def stop(self):
        if self.proc and self.proc.poll() is None:
            self.proc.kill()
        self.proc = None


worker = _Worker()                           # the page's viewport, renders, exports (interactive)
bg = _Worker()                               # long analyses (frames, pile-up, regime, research): never block the viewport


def _prune():
    files = sorted(OUT_DIR.glob("*"), key=lambda p: p.stat().st_mtime)
    for p in files[:-KEEP_FILES]:
        try:
            p.unlink()
        except OSError:
            pass


def _source(body):
    """{path, pattern?} or {upload, pattern} -> OVITO import path (wildcard sequence or one file scanned for all its frames)."""
    if body.get("remote"):                       # HPC files: streamed into memory, never written to this computer
        pat = str(body.get("pattern") or "")
        if not pat or "/" in pat or not body.get("dir"):
            raise HTTPException(422, "remote folder and file pattern required")
        return dict(import_path=remote.stream_url(body), multiple_frames="*" not in pat, label=body.get("label") or remote.label(body),
                    pattern=pat, remote=str(body["remote"]), dir=str(body["dir"]))
    if body.get("blob"):                         # hosted uploads, fetched frame by frame (remote.py)
        folder = remote.mirror(body)[0]
        pat = str(body["pattern"])
        extra = dict(remote=str(body["remote"]), dir=str(body["dir"])) if body.get("remote") else dict(blob=body["blob"])
        return dict(import_path=str(folder / pat), multiple_frames="*" not in pat, label=body.get("label") or remote.label(body), pattern=pat, **extra)
    if body.get("upload"):                       # files uploaded from the browser (this server's temporary folder)
        sid = str(body["upload"])
        if not re.fullmatch(r"[0-9a-f]{16}", sid) or not (UPLOAD_DIR / sid).is_dir():
            raise HTTPException(404, "upload not found (uploads are kept for 24 hours); upload the files again")
        p = UPLOAD_DIR / sid
    else:
        try:
            p = studio.resolve(body.get("path"))
        except Exception as e:
            raise HTTPException(404 if isinstance(e, FileNotFoundError) else 403 if isinstance(e, PermissionError) else 422, str(e))
    pattern = body.get("pattern")
    if p.is_file():
        if not studio.is_dump(p):
            raise HTTPException(422, "not a LAMMPS text dump")
        return dict(import_path=str(p), multiple_frames=True, label=p.name)
    if not pattern or "/" in pattern or "\\" in pattern:
        raise HTTPException(422, "pattern (file name or wildcard) required for a folder")
    full = p / pattern
    single = "*" not in pattern
    return dict(import_path=str(full), multiple_frames=single, label=f"{p.name}/{pattern}")


@router.get("/status")
def status():
    """Starts the OVITO worker if needed and reports the OVITO version (the first call takes a few seconds)."""
    try:
        info = worker.ask({"op": "hello"}, 120)
        return dict(available=True, ovito=info["ovito"], python=sys.version.split()[0], hosted=HOSTED, uploads="blob" if HOSTED else "server")
    except Exception as e:
        return dict(available=False, error=str(e), install=r".venv\Scripts\python -m pip install ovito")


@router.post("/info")
def info(body: dict = Body(...)):
    src = _source(body)
    frame = int(body.get("frame", 0))
    key = f"ovito-info:{src['import_path']}:{frame}:{body.get('refresh', 0)}"
    def run(progress=None):
        _prep(src, [frame], progress)
        return dict(source=src, **worker.ask({"op": "info", "frame": frame, **src}))
    return _job(key, "OVITO reads the trajectory", run)


def _out_url(path, ctype="application/octet-stream"):
    """Where the browser downloads a result file: this server's /api/ovito/file locally; a presigned Blob URL when
    hosted (the next request may reach another instance, and responses are limited to 4.5 MB)."""
    path = Path(path)
    if HOSTED and blobstore.enabled():
        url = blobstore.publish(str(path), path.name, ctype)
        path.unlink(missing_ok=True)
        return url
    return f"/api/ovito/file/{path.name}"


def _refs(spec):
    """Reference frames the pipeline reads besides the current one (Displacement vectors)."""
    return [int(m.get("reference_frame", 0)) for m in spec.get("modifiers") or [] if m.get("enabled", True) and m.get("type") == "displacements"]


def _prep(src, frames, progress=None):
    """Remote sources: fetch the files of these frames from the cluster first (no-op for local files)."""
    remote.prepare(src, frames, progress)


def _spec(body):
    spec = body.get("spec") or {}
    src = _source(spec.get("source") or {})
    spec = dict(spec, source=src)
    ovito_script.build(spec)                     # validates every value; raises ValueError with a readable message
    return spec


@router.post("/script")
def script(body: dict = Body(...)):
    """The OVITO Python script for the current settings (as run by the worker, with import_file instead of the cache)."""
    try:
        spec = _spec(body)
        return dict(script=ovito_script.build(spec, image_path="nanocut_ovito.png"))
    except ValueError as e:
        raise HTTPException(422, str(e))


@router.post("/render")
def render(body: dict = Body(...)):
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    spec = dict(spec, render=dict(spec.get("render") or {}, anim=None))   # one frame per request (movies: the page loops)
    image = OUT_DIR / f"render_{uuid.uuid4().hex[:16]}.png"

    def run(progress=None):
        _prep(spec["source"], [int(spec.get("frame", 0))] + _refs(spec), progress)
        res = worker.ask({"op": "run", "spec": spec, "image_path": str(image)})
        size = image.stat().st_size
        _prune()
        return dict(url=_out_url(image, "image/png"), bytes=size, script=ovito_script.build(spec, image_path="nanocut_ovito.png"), **res)
    return _job(f"ovito-render:{image.name}", "OVITO renders the image", run)


@router.post("/export")
def export(body: dict = Body(...)):
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    fmt = body.get("format", "lammps/dump")
    if fmt not in EXPORT_FORMATS:
        raise HTTPException(422, f"format must be one of {', '.join(EXPORT_FORMATS)}")
    cols = [c for c in body.get("columns") or [] if isinstance(c, str) and ovito_script.PROPERTY_RE.match(c)]
    if fmt in ("lammps/dump", "xyz") and not cols:
        raise HTTPException(422, "choose at least one column to export")
    name = f"export_{uuid.uuid4().hex[:10]}{EXPORT_FORMATS[fmt]}"

    def run(progress=None):
        _prep(spec["source"], [int(spec.get("frame", 0))] + _refs(spec), progress)
        res = worker.ask({"op": "export", "spec": spec, "path": str(OUT_DIR / name), "format": fmt, "columns": cols})
        _prune()
        return dict(file=name, bytes=res["bytes"], url=_out_url(OUT_DIR / name))
    return _job(f"ovito-export:{name}", "OVITO exports the data", run)


@router.post("/verify")
def verify(body: dict = Body(...)):
    """Cross-check: OVITO's Displacement vectors vs NanoCut's own dump parser (the Studio) for one frame."""
    src = _source(body)
    if src.get("remote"):
        raise HTTPException(422, "the cross-check reads files from disk; HPC trajectories are only streamed into memory")
    frame, ref = int(body.get("frame", 0)), int(body.get("reference", 0))

    def run(progress=None):
        _prep(src, [frame, ref], progress)
        npz = OUT_DIR / f"verify_{uuid.uuid4().hex[:8]}.npz"
        o = worker.ask({"op": "displacements", "frame": frame, "reference": ref, "out": str(npz), **src})
        with np.load(npz) as f:
            z = {k: f[k] for k in f.files}          # copy, so the file can be deleted (Windows)
        directory, pattern = (str(Path(src["import_path"]).parent), Path(src["import_path"]).name)
        s = studio.open_source(directory, pattern)
        meta = studio.load_frame(s["id"], frame, ref)
        n = meta["natoms"]
        raw = np.frombuffer(studio.columns_bytes(s["id"], frame, ["id", "x", "y", "z", "dx", "dy", "dz"], ref), dtype="<f4")
        ids = np.frombuffer(raw[:n].tobytes(), "<u4").astype(np.int64)
        cols = raw[n:].reshape(6, n)
        npz.unlink(missing_ok=True)
        same_ids = len(ids) == len(z["ids"]) and np.array_equal(ids, z["ids"])
        out = dict(frame=frame, reference=ref, step_ovito=o["step"], step_nanocut=meta["step"], atoms_ovito=int(len(z["ids"])),
                   atoms_nanocut=int(n), same_atom_ids=bool(same_ids))
        if same_ids:
            for k, name in enumerate(("x", "y", "z")):
                out[f"max_diff_position_{name}_A"] = float(np.max(np.abs(cols[k] - z["pos"][:, k])))
            for k, name in enumerate(("dx", "dy", "dz")):
                out[f"max_diff_{name}_A"] = float(np.nanmax(np.abs(cols[3 + k] - z["disp"][:, k])))
                out[f"range_{name}_ovito"] = [float(z["disp"][:, k].min()), float(z["disp"][:, k].max())]
        return out
    return _job(f"ovito-verify:{src['import_path']}:{frame}:{ref}:{time.time():.0f}", "OVITO vs NanoCut cross-check", run)


# ---------------------------------------------------------------------------------------------- live 3D viewport
@router.post("/live")
def live(body: dict = Body(...)):
    """OVITO evaluates the page's pipeline for one frame; the browser draws the result in 3D. Response: uint32 length
    of a JSON header (OVITO's attributes, statistics, legend, layout), the header, then per particle after the
    pipeline position, display colour and radius (layout "f32": float32 x3, uint8 x3 padded, float32; layout "q16"
    when hosted frames would exceed the response limit, see ovito_worker._encode), then dislocation line vertices."""
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    out = OUT_DIR / f"live_{uuid.uuid4().hex[:12]}.bin"
    try:
        _prep(spec["source"], [int(spec.get("frame", 0))] + _refs(spec))
        meta = worker.ask({"op": "live", "spec": spec, "frame": spec.get("frame", 0), "out": str(out),
                           "max_bytes": int(os.environ.get("NANOCUT_LIVE_MAX_BYTES", "3600000")) if HOSTED else 0})   # hosted: < 4.5 MB
        data = out.read_bytes()
    except RuntimeError as e:
        raise HTTPException(500, str(e))
    finally:
        out.unlink(missing_ok=True)
    head = json.dumps(meta).encode()
    head += b" " * ((-len(head) - 4) % 4)                    # binary part starts 4-byte aligned
    return Response(content=len(head).to_bytes(4, "little") + head + data, media_type="application/octet-stream")


@router.post("/atom")
def atom(body: dict = Body(...)):
    """All OVITO particle properties of one particle (index after the pipeline) of a live frame: {spec, index}."""
    try:
        spec = _spec(body)
        _prep(spec["source"], [int(spec.get("frame", 0))] + _refs(spec))
        return worker.ask({"op": "atom", "spec": spec, "frame": int(spec.get("frame", 0)), "index": int(body.get("index", -1))}, 600)
    except ValueError as e:
        raise HTTPException(422, str(e))
    except RuntimeError as e:
        raise HTTPException(409, str(e))


# ---------------------------------------------------------------------------------------------- frame export & statistics
_cancel = set()
MAX_FRAMES = 100_000


def _dump_columns(src):
    """Column names of the first dump file of the source (as written by LAMMPS, from its 'ITEM: ATOMS' line)."""
    if src.get("remote"):
        return remote.dump_columns(src)
    p = Path(src["import_path"])
    files = sorted(p.parent.glob(p.name), key=lambda f: [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", f.name)]) if "*" in p.name else [p]
    if not files:
        return []
    try:
        return studio._header_at(str(files[0]), 0)["columns"]
    except Exception:
        return []


@router.post("/frames")
def frames_job(body: dict = Body(...)):
    """Selected frames through OVITO: a ZIP with one CSV per frame (all particle columns) and/or per-frame statistics
    (count, mean, std, min, max of every numeric column; force averages). body: {spec, frames, apply_pipeline,
    include_derived, csv, stats, token}."""
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    want_csv, want_stats = bool(body.get("csv")), bool(body.get("stats"))
    if not (want_csv or want_stats):
        raise HTTPException(422, "nothing to do: ask for csv and/or stats")
    try:
        frames = sorted({int(f) for f in body.get("frames") or []})
    except (TypeError, ValueError):
        raise HTTPException(422, "frames must be integers")
    if not frames:
        raise HTTPException(422, "no frames selected")
    if len(frames) > MAX_FRAMES:
        raise HTTPException(422, f"at most {MAX_FRAMES} frames per export")
    token = str(body.get("token") or uuid.uuid4().hex)
    src = spec["source"]
    save_to = None
    if body.get("save_to") and want_csv and not HOSTED:       # a path from the local Save-as dialog
        save_to = Path(str(body["save_to"]))
        if not save_to.is_absolute() or save_to.suffix.lower() != ".zip" or not save_to.parent.is_dir():
            raise HTTPException(422, "save_to must be an absolute path of a .zip file in an existing folder")
        save_to.unlink(missing_ok=True)                        # the dialog already asked before replacing it

    def run(progress=None):
        _prep(src, [0], progress)
        info = bg.ask({"op": "info", "frame": 0, **src})
        n = info["num_frames"]
        bad = [f for f in frames if not 0 <= f < n]
        if bad:
            raise ValueError(f"frames outside 0 … {n - 1}: {bad[:20]}{' …' if len(bad) > 20 else ''}")
        dump_cols = _dump_columns(src)
        zip_name = f"frames_{uuid.uuid4().hex[:10]}.zip" if want_csv and not HOSTED else None
        zip_file = save_to or (OUT_DIR / zip_name if zip_name else None)     # local: straight to the chosen file
        rows, t0 = [], time.time()
        try:
            for k, f in enumerate(frames):
                if token in _cancel:
                    raise RuntimeError(f"cancelled after {k} of {len(frames)} frames")
                if progress:
                    eta = (time.time() - t0) / k * (len(frames) - k) if k else None
                    progress(k / len(frames), f"frame {f} ({k + 1} of {len(frames)})" + (f", about {eta:.0f} s left" if eta else ""))
                _prep(src, [f] + (_refs(spec) if body.get("apply_pipeline") else []))
                csv_path = OUT_DIR / f"frame_{f}_{uuid.uuid4().hex[:8]}.csv" if want_csv and HOSTED else None
                row = bg.ask({"op": "frame_table", "spec": spec, "frame": f, "apply_pipeline": bool(body.get("apply_pipeline")),
                                  "include_derived": bool(body.get("include_derived")), "dump_columns": dump_cols,
                                  "raw_properties": info["properties"], "stats": want_stats,
                                  "zip_path": str(zip_file) if zip_file else None, "csv_path": str(csv_path) if csv_path else None})
                if csv_path:
                    row["csv_bytes"] = csv_path.stat().st_size
                    row["csv_url"] = _out_url(csv_path, "text/csv")
                rows.append(row)
        except Exception:
            if zip_file:
                Path(zip_file).unlink(missing_ok=True)
            raise
        finally:
            _cancel.discard(token)
        _prune()
        meta = dict(source=src.get("label"), import_path=src["import_path"], total_frames=n, frames=frames, dump_columns=dump_cols,
                    apply_pipeline=bool(body.get("apply_pipeline")), include_derived=bool(body.get("include_derived")),
                    modifiers=[m.get("type") for m in spec.get("modifiers") or [] if m.get("enabled", True)] if body.get("apply_pipeline") else [],
                    ovito=bg.info.get("ovito") if bg.info else None, seconds=round(time.time() - t0, 2))
        return dict(zip=zip_name, url=None if save_to or not zip_name else f"/api/ovito/file/{zip_name}", saved_to=str(save_to) if save_to else None,
                    bytes=Path(zip_file).stat().st_size if zip_file else sum(r.get("csv_bytes", 0) for r in rows), rows=rows, meta=meta)
    key = hashlib.sha1(json.dumps([frames, body.get("apply_pipeline"), body.get("include_derived"), want_csv, want_stats]).encode()).hexdigest()[:12]
    return _job(f"ovito-frames:{token}:{key}", "OVITO exports frames" if want_csv else "OVITO computes frame statistics", run)


REGIME_PARAMS = {"probe_radius": (1.0, 10.0), "ahead": (0.0, 2000.0), "behind": (0.0, 2000.0), "crack_margin": (0.0, 50.0),
                 "bottom_margin": (0.0, 200.0), "side_margin": (0.0, 500.0), "void_min_volume": (0.0, 1e6), "chip_threshold": (0.0, 50.0), "cluster_cutoff": (1.0, 10.0),
                 "min_fragment": (1, 1e7)}


@router.post("/pileup")
def pileup_job(body: dict = Body(...)):
    """Pile-up height per selected frame (ovito_worker.op_pileup). body: {spec, frames, params: {tool_types, threshold, direction}, token}."""
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    try:
        frames = sorted({int(f) for f in body.get("frames") or []})
    except (TypeError, ValueError):
        raise HTTPException(422, "frames must be integers")
    if not frames or len(frames) > MAX_FRAMES:
        raise HTTPException(422, "select between 1 and 100000 frames")
    raw = body.get("params") or {}
    prm = {"tool_types": [int(t) for t in raw.get("tool_types") or []], "direction": "+x" if raw.get("direction") == "+x" else "-x",
           "threshold": float(raw.get("threshold", 1.0))}
    if not 0 <= prm["threshold"] <= 100:
        raise HTTPException(422, "threshold must be within 0 … 100 Å")
    token, src = str(body.get("token") or uuid.uuid4().hex), spec["source"]

    def run(progress=None):
        rows, t0 = [], time.time()
        try:
            for k, f in enumerate(frames):
                if token in _cancel:
                    raise RuntimeError(f"cancelled after {k} of {len(frames)} frames")
                if progress:
                    progress(k / len(frames), f"frame {f} ({k + 1} of {len(frames)})")
                _prep(src, [f])
                rows.append(bg.ask({"op": "pileup", "spec": spec, "frame": f, "params": prm}))
        finally:
            _cancel.discard(token)
        return dict(rows=rows, params=prm, seconds=round(time.time() - t0, 2))
    key = hashlib.sha1(json.dumps([frames, prm]).encode()).hexdigest()[:12]
    return _job(f"ovito-pileup:{token}:{key}", "OVITO measures the pile-up", run)


@router.post("/regime")
def regime_job(body: dict = Body(...)):
    """Ductile/brittle indicators for the selected frames (ovito_worker.op_regime). body: {spec, frames, params, token}."""
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    try:
        frames = sorted({int(f) for f in body.get("frames") or []})
    except (TypeError, ValueError):
        raise HTTPException(422, "frames must be integers")
    if not frames or len(frames) > MAX_FRAMES:
        raise HTTPException(422, "select between 1 and 100000 frames")
    raw = body.get("params") or {}
    prm = {"tool_types": [int(t) for t in raw.get("tool_types") or []], "direction": "+x" if raw.get("direction") == "+x" else "-x"}
    if not prm["tool_types"]:
        raise HTTPException(422, "choose the particle type(s) of the tool")
    for k, (lo, hi) in REGIME_PARAMS.items():
        if raw.get(k) is not None:
            v = float(raw[k])
            if not lo <= v <= hi:
                raise HTTPException(422, f"{k} must be within {lo} … {hi}")
            prm[k] = v
    token, src = str(body.get("token") or uuid.uuid4().hex), spec["source"]

    def run(progress=None):
        rows, t0 = [], time.time()
        try:
            for k, f in enumerate(frames):
                if token in _cancel:
                    raise RuntimeError(f"cancelled after {k} of {len(frames)} frames")
                if progress:
                    eta = (time.time() - t0) / k * (len(frames) - k) if k else None
                    progress(k / len(frames), f"frame {f} ({k + 1} of {len(frames)})" + (f", about {eta:.0f} s left" if eta else ""))
                _prep(src, [f])
                rows.append(bg.ask({"op": "regime", "spec": spec, "frame": f, "params": prm}))
        finally:
            _cancel.discard(token)
        return dict(rows=rows, params=prm, seconds=round(time.time() - t0, 2), source=src.get("label"))
    key = hashlib.sha1(json.dumps([frames, prm]).encode()).hexdigest()[:12]
    return _job(f"ovito-regime:{token}:{key}", "OVITO computes ductile/brittle indicators", run)


def _find_params(src):
    """params.json of a NanoCut run next to the dumps (or one or two folders up), for local sources."""
    if src.get("remote") or src.get("blob"):
        return None
    d = Path(src["import_path"]).parent
    for p in (d, d.parent, d.parent.parent):
        f = p / "params.json"
        if f.is_file():
            try:
                return f, json.loads(f.read_text(encoding="utf-8"))
            except Exception:
                return None
    return None


@router.post("/research/defaults")
def research_defaults(body: dict = Body(...)):
    """Suggested analysis settings: types (name, mass, atomic volume, role), tool types, time step, structure method.
    From the run's params.json when found, else from the type names (element masses; volumes to be entered)."""
    from nanocut import research
    from nanocut.materials import MASS, MATERIALS
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    info_ = bg.ask({"op": "info", "frame": 0, **spec["source"]})
    out = dict(types={}, tool_types=[], dt_ps=None, structure="none", source="type names", direction="-x")
    found = _find_params(spec["source"])
    if found:
        f, prm = found
        exp = prm.get("experiment", prm)
        try:
            out["types"] = {str(k): v for k, v in research.type_table_from_experiment(exp).items()}
            out["tool_types"] = [1]
            out["dt_ps"] = float(exp.get("thermo", {}).get("cut_dt_fs", 0.5)) * 1e-3
            mat = MATERIALS.get(exp["workpiece"]["material"], {})
            out["structure"] = "diamond" if mat.get("structure") in ("zincblende", "diamond") else "cna" if mat.get("structure") in ("fcc", "bcc") else "ptm"
            out["source"] = str(f)
            out["rake_deg"] = float(exp.get("tool", {}).get("rake_deg", 0.0))
        except Exception as e:                                   # an older/foreign params.json: fall back to names
            out["note"] = f"params.json not usable ({e})"
    for t in info_["types"]:
        k = str(t["id"])
        if k not in out["types"]:
            out["types"][k] = dict(name=t.get("name") or "", mass=MASS.get(t.get("name") or ""), volume=None, role="workpiece")
        out["types"][k]["count"] = t["count"]
    return out


@router.post("/research/sidefiles")
def research_sidefiles(body: dict = Body(...)):
    """LAMMPS text output next to the dumps (fix ave/time files such as *_forces_temps.txt, log files), parsed into
    tables for Graph Studio. Local sources only; files above 50 MB and dump files are skipped."""
    from . import graph_import
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    src = spec["source"]
    if src.get("remote") or src.get("blob") or src.get("upload"):
        return dict(datasets=[])
    d = Path(src["import_path"]).parent
    out, seen = [], set()
    for folder in (d, d.parent):
        for f in sorted(folder.glob("*")):
            if f in seen or not f.is_file() or f.stat().st_size > 50 << 20 or not (f.suffix.lower() in (".txt", ".log", ".dat", ".csv") or f.name.startswith("log")):
                continue
            seen.add(f)
            try:
                head = f.read_bytes()[:4000]
                if b"ITEM: TIMESTEP" in head:
                    continue
                for ds in graph_import.parse(f.name, f.read_bytes()):
                    if len(ds.get("columns") or {}) >= 2:
                        out.append(dict(ds, path=str(f)))
            except Exception:
                continue
    return dict(datasets=out[:20])


@router.post("/research/volumes")
def research_volumes(body: dict = Body(...)):
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    def run(progress=None):
        _prep(spec["source"], [0], progress)
        return bg.ask({"op": "volumes", "spec": spec})
    return _job(f"ovito-volumes:{spec['source']['import_path']}", "OVITO estimates atomic volumes (Voronoi, frame 0)", run)


@router.post("/research")
def research_job(body: dict = Body(...)):
    """Research analysis of the selected frames (PLAN_PHASE3.md §2). body: {spec, frames, cfg, token}."""
    from nanocut import research
    try:
        spec = _spec(body)
    except ValueError as e:
        raise HTTPException(422, str(e))
    try:
        frames = sorted({int(f) for f in body.get("frames") or []})
    except (TypeError, ValueError):
        raise HTTPException(422, "frames must be integers")
    if not frames or len(frames) > MAX_FRAMES:
        raise HTTPException(422, "select between 1 and 100000 frames")
    cfg = dict(body.get("cfg") or {})
    for k, (lo, hi) in {"contact_radius": (1, 500), "chip_threshold": (0, 50), "subsurface_depth": (1, 500), "rdf_rmax": (2, 20),
                        "rdf_bins": (20, 1000), "max_centers": (100, 200000), "map_bin": (0.5, 50), "profile_bin": (0.2, 50),
                        "hist_bins": (5, 500), "bad_bins": (10, 720), "strain_cutoff": (1.5, 10), "cn_cutoff": (0.5, 6)}.items():
        if cfg.get(k) is not None and not lo <= float(cfg[k]) <= hi:
            raise HTTPException(422, f"{k} must be within {lo} … {hi}")
    if cfg.get("structure") not in (None, "none", "diamond", "cna", "ptm"):
        raise HTTPException(422, "structure must be none, diamond, cna or ptm")
    token, src = str(body.get("token") or uuid.uuid4().hex), spec["source"]

    def run(progress=None):
        rows, t0 = [], time.time()
        try:
            _prep(src, [0])
            for k, f in enumerate(frames):
                if token in _cancel:
                    raise RuntimeError(f"cancelled after {k} of {len(frames)} frames")
                if progress:
                    eta = (time.time() - t0) / k * (len(frames) - k) if k else None
                    progress(k / len(frames), f"frame {f} ({k + 1} of {len(frames)})" + (f", about {eta:.0f} s left" if eta else ""))
                _prep(src, [f])
                rows.append(bg.ask({"op": "research", "spec": spec, "frame": f, "cfg": cfg}))
        finally:
            _cancel.discard(token)
        rows = research.finalize(rows, cfg)
        return dict(rows=rows, catalogue=research.catalogue(rows), cfg=cfg, source=src.get("label"), import_path=src["import_path"],
                    seconds=round(time.time() - t0, 2), ovito=bg.info.get("ovito") if bg.info else None,
                    params=_find_params(src)[1] if _find_params(src) else None)
    key = hashlib.sha1(json.dumps([frames, cfg], sort_keys=True).encode()).hexdigest()[:12]
    return _job(f"ovito-research:{token}:{key}", "OVITO research analysis", run)


@router.post("/frames/cancel/{token}")
def frames_cancel(token: str):
    _cancel.add(token)
    return dict(cancelling=token)


@router.post("/frames/xlsx")
def frames_xlsx(body: dict = Body(...)):
    """Excel workbook of the frame statistics: Frame Statistics (force averages), Export Metadata, then one sheet per
    group of columns (computes c_*, fixes f_*, variables v_*, displacement, positions/velocities, others), restricted
    to the columns and statistics chosen on the page. body: {rows, meta, columns: [...], stats: [...], selection: {...}}."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    rows, meta = body.get("rows") or [], body.get("meta") or {}
    label_by = {"n": "valid count", "mean": "mean", "std": "std (population)", "min": "min", "max": "max"}
    chosen = [c for c in body.get("columns") or [] if isinstance(c, str)]
    stat_keys = [s for s in body.get("stats") or ["mean"] if s in ("n", "mean", "std", "min", "max")] or ["mean"]
    if not rows:
        raise HTTPException(422, "no statistics to export; calculate them first")
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F3B4A")

    def sheet(ws, header, data, widths=None):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in data:
            ws.append(r)
        ws.freeze_panes = "B2"
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(12, min(40, len(str(h)) + 2)))
        for col in ws.iter_cols(min_row=2):
            for c in col:
                if isinstance(c.value, float):
                    c.number_format = "0.000000E+00" if c.value and (abs(c.value) >= 1e6 or abs(c.value) < 1e-3) else "0.000000"

    ws = wb.active
    ws.title = "Frame Statistics"
    FK = ("mean_fx", "mean_fy", "mean_fz", "mean_magnitude", "magnitude_of_mean", "sum_fx", "sum_fy", "sum_fz")
    FL = ["Mean Fx", "Mean Fy", "Mean Fz", "Mean force magnitude  mean(|F_i|)", "Magnitude of mean force  |mean(F)|", "Sum Fx", "Sum Fy", "Sum Fz"]
    sheet(ws, ["Frame", "Timestep", "Particle count", "Valid force particles"] + FL,
          [[r["frame"], r.get("timestep"), r.get("count"), (r.get("force") or {}).get("n_valid"), *[(r.get("force") or {}).get(k) for k in FK]] for r in rows])
    types = sorted({t for r in rows for t in (r.get("force_by_type") or {})}, key=int)
    by_type_ws = None
    if types and any((r.get("force") or {}).get("n_valid") is not None for r in rows):
        by_type_ws = (types, [[r["frame"], r.get("timestep")] + [((r.get("force_by_type") or {}).get(t) or {}).get(k) for t in types for k in ("n_valid",) + FK] for r in rows])
    groups = [("Computes (c_)", lambda c: c.startswith("c_")), ("Fixes (f_)", lambda c: c.startswith("f_")), ("Variables (v_)", lambda c: c.startswith("v_")),
              ("Displacement", lambda c: c.startswith("Displacement")),
              ("Positions, velocities, forces", lambda c: c in ("x", "y", "z", "xu", "yu", "zu", "vx", "vy", "vz", "fx", "fy", "fz") or c.split(".")[0] in ("Position", "Velocity", "Force", "Velocity Magnitude")),
              ("Other properties", lambda c: True)]
    meta_ws = wb.create_sheet("Export Metadata")             # sheet 2; filled at the end
    tkeys = sorted({t for r in rows for t in (r.get("stats_by_type") or {})}, key=int)
    if tkeys and chosen:                                       # every chosen column per particle type, one row per frame and type
        sheet(wb.create_sheet("Columns by particle type"), ["Frame", "Timestep", "Particle type", "Particles"] + [f"{c} {label_by[s_]}" for c in chosen for s_ in stat_keys],
              [[r["frame"], r.get("timestep"), int(t), ((r.get("stats_by_type") or {}).get(t, {}).get(chosen[0]) or {}).get("n")] +
               [(((r.get("stats_by_type") or {}).get(t) or {}).get(c) or {}).get(s_) for c in chosen for s_ in stat_keys]
               for r in rows for t in tkeys if t in (r.get("stats_by_type") or {})])
    if by_type_ws:
        tnames = {t: next(((r.get("force_by_type") or {}).get(t, {}).get("name") for r in rows if (r.get("force_by_type") or {}).get(t)), "") for t in by_type_ws[0]}
        sheet(wb.create_sheet("Force by particle type"), ["Frame", "Timestep"] +
              [f"type {t}{' ' + tnames[t] if tnames[t] else ''}: {h}" for t in by_type_ws[0] for h in ["valid particles"] + FL], by_type_ws[1])
    placed = set()
    label = {"n": "valid count", "mean": "mean", "std": "std (population)", "min": "min", "max": "max"}
    for title, test in groups:
        cols = [c for c in chosen if c not in placed and test(c)]
        if not cols:
            continue
        placed.update(cols)
        header = ["Frame", "Timestep"] + [f"{c} {label[s]}" for c in cols for s in stat_keys]
        data = [[r["frame"], r.get("timestep")] + [((r.get("stats") or {}).get(c) or {}).get(s) for c in cols for s in stat_keys] for r in rows]
        sheet(wb.create_sheet(title[:31]), header, data)
    sel = body.get("selection") or {}
    origins = {}
    for r in rows:
        for c in r.get("columns") or []:
            origins.setdefault(c["name"], f"{c['origin']} (OVITO: {c['ovito']})")
    md = [["Source", meta.get("source")], ["Import path", meta.get("import_path")], ["Total frames available", meta.get("total_frames")],
          ["Selected frames", ", ".join(map(str, meta.get("frames") or []))], ["Number of selected frames", len(meta.get("frames") or [])],
          ["Frame-selection mode", sel.get("mode")], ["Selection parameters", json.dumps({k: v for k, v in sel.items() if k != "mode"})],
          ["Data", "after the OVITO pipeline: " + ", ".join(meta.get("modifiers") or []) if meta.get("apply_pipeline") else "as read from the dump file (no modifiers)"],
          ["Columns derived by OVITO included", "yes" if meta.get("include_derived") else "no (only dump columns and pipeline outputs)"],
          ["Force properties used", "fx, fy, fz (OVITO Force.X, Force.Y, Force.Z)" if any((r.get("force") or {}).get("n_valid") is not None for r in rows) else
           "none: the dump has no complete force vector (" + ", ".join(((rows[0].get("force") or {}).get("missing") or [])) + " missing)"],
          ["Force averages", "N = particles whose Fx, Fy and Fz are all finite; Mean Fx = sum(Fx_i)/N (likewise y, z); "
                             "mean |F| = sum(sqrt(Fx_i^2+Fy_i^2+Fz_i^2))/N; |mean F| = sqrt(MeanFx^2+MeanFy^2+MeanFz^2); Sum Fx = sum(Fx_i) over the N particles."],
          ["Note on forces", "Over all atoms of an isolated system the internal forces cancel (Newton's third law), so the all-atom mean and sum are close to 0; "
                             "the force on the tool is the sum over the tool's particle type (sheet 'Force by particle type'). Units: those of the dump (LAMMPS metal: eV/Å)."],
          ["Column statistics", "over the particles whose value is finite (valid count); std is the population standard deviation (ddof = 0)."],
          ["Statistics exported", ", ".join(label[s] for s in stat_keys)], ["OVITO version", meta.get("ovito")],
          ["Exported", time.strftime("%Y-%m-%d %H:%M:%S %Z")], ["Columns (origin)", ""]] + [[k, v] for k, v in origins.items()]
    sheet(meta_ws, ["Item", "Value"], md, widths={1: 34, 2: 110})
    buf = BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------------------------------------- uploads
_last_cleanup = [0.0]
MAX_UPLOAD = int(os.environ.get("NANOCUT_MAX_UPLOAD", str(2 << 30)))


@router.post("/upload-token")
def upload_token(body: dict = Body(...)):
    """Hosted: a client token that lets the browser upload one file straight to the private Blob store (nothing passes
    through this function, so the 4.5 MB request limit does not apply). Returns where and how to upload it."""
    if not (HOSTED and blobstore.enabled()):
        raise HTTPException(404, "uploads go to this server (POST /api/ovito/upload/...)")
    name, size = str(body.get("name") or ""), int(body.get("size") or 0)
    if not re.fullmatch(r"[\w.\-]{1,200}", name) or name.startswith("."):
        raise HTTPException(422, "invalid file name")
    if not 0 < size <= MAX_UPLOAD:
        raise HTTPException(413, f"files up to {MAX_UPLOAD / 2**30:.0f} GB can be uploaded to the hosted version")
    if time.time() - _last_cleanup[0] > 3600:                # uploads and results are kept for a day
        _last_cleanup[0] = time.time()
        threading.Thread(target=blobstore.cleanup, daemon=True).start()
    session = re.sub(r"[^0-9a-f]", "", str(body.get("session") or ""))[:16] or uuid.uuid4().hex[:16]
    pathname = f"{blobstore.UPLOAD_PREFIX}{time.strftime('%Y%m%d')}/{session}/{name}"
    return dict(pathname=pathname, token=blobstore.client_token(pathname, size), store=blobstore.store_id(), api=blobstore.API,
                version=blobstore.VERSION, multipart=size > 64 << 20, part_size=8 << 20, session=session)
@router.post("/upload/{sid}/{name}")
async def upload(sid: str, name: str, request: Request, offset: int = 0):
    """Append one chunk of a file uploaded from the browser to this server's temporary folder (kept 24 hours).
    offset must equal the bytes already received, so an interrupted upload can continue where it stopped."""
    if not re.fullmatch(r"[0-9a-f]{16}", sid) or not re.fullmatch(r"[\w.\-]{1,200}", name) or name.startswith("."):
        raise HTTPException(422, "invalid upload id or file name")
    for d in UPLOAD_DIR.iterdir():
        if d.is_dir() and d.name != sid and time.time() - d.stat().st_mtime > UPLOAD_KEEP_S:
            shutil.rmtree(d, ignore_errors=True)
    folder = UPLOAD_DIR / sid
    folder.mkdir(exist_ok=True)
    path = folder / name
    have = path.stat().st_size if path.exists() else 0
    if offset != have:
        raise HTTPException(409, f"expected offset {have}")
    chunk = await request.body()
    with open(path, "ab") as f:
        f.write(chunk)
    os.utime(folder)
    return dict(name=name, bytes=have + len(chunk))


@router.get("/movie/{folder}/{name}")
def movie_frame(folder: str, name: str):
    p = OUT_DIR / folder / name
    if not re.fullmatch(r"movie_[0-9a-f]{16}", folder) or not re.fullmatch(r"frame_\d+\.png", name) or not p.is_file():
        raise HTTPException(404, "movie frame not found (the latest 3 movies are kept)")
    return FileResponse(p, media_type="image/png")


@router.get("/file/{name}")
def file(name: str):
    p = (OUT_DIR / name).resolve()
    if p.parent != OUT_DIR.resolve() or not p.is_file():
        raise HTTPException(404, "file not found (renders are kept for the latest 60 results)")
    media = {".png": "image/png", ".mp4": "video/mp4"}.get(p.suffix, "application/octet-stream")
    return FileResponse(p, media_type=media, filename=p.name)


def _job(key, label, fn):
    from fastapi.responses import JSONResponse
    rec = jobs.submit(key, label, fn)
    return JSONResponse(jobs.public(rec), status_code=200 if rec["status"] == "done" else 202)
